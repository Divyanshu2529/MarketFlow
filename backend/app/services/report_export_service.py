from __future__ import annotations

import json
import re
from datetime import date, datetime
from io import BytesIO
from typing import Any

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import (
    ParagraphStyle,
    getSampleStyleSheet,
)
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from xml.sax.saxutils import escape

from app.models.report import Report


REPORT_SECTIONS = [
    ("Company Overview", "company"),
    ("Financial History", "history"),
    ("Calculated P/E Ratio", "calculatedPeRatio"),
    ("AI Recommendation", "recommendation"),
    ("Sentiment Analysis", "sentiment"),
    ("News", "news"),
    ("SEC Filings", "filings"),
    ("Competitor Comparison", "competitors"),
]

LONG_TEXT_KEYS = {
    "description",
    "summary",
    "analysis",
    "explanation",
    "reason",
    "content",
    "text",
}


def humanize_key(key: Any) -> str:
    text = str(key).replace("_", " ").replace("-", " ")
    text = re.sub(r"(?<!^)(?=[A-Z])", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text.title()


def _format_date(value: str) -> str:
    try:
        parsed = datetime.fromisoformat(
            value.replace("Z", "+00:00")
        )
        return parsed.strftime("%B %d, %Y")
    except ValueError:
        return value


def format_value(key: Any, value: Any) -> str:
    if value is None:
        return "Not available"

    if isinstance(value, bool):
        return "Yes" if value else "No"

    if isinstance(value, (datetime, date)):
        return value.strftime("%B %d, %Y")

    if isinstance(value, str):
        lower_key = str(key).lower()

        if (
            "date" in lower_key
            or lower_key.endswith("at")
        ):
            return _format_date(value)

        return value

    if isinstance(value, (int, float)):
        lower_key = str(key).lower()

        if (
            "percentage" in lower_key
            or lower_key.endswith("percent")
        ):
            return f"{value:,.2f}%"

        if "margin" in lower_key:
            percentage = value * 100 if abs(value) <= 1 else value
            return f"{percentage:,.2f}%"

        if lower_key in {"price", "eps"}:
            return f"${value:,.2f}"

        if lower_key == "peratio":
            return f"{value:,.2f}"

        if any(
            term in lower_key
            for term in (
                "marketcap",
                "revenue",
                "cashflow",
                "dividend",
                "assets",
                "liabilities",
                "income",
                "sales",
            )
        ):
            return f"${value:,.2f}"

        if lower_key in {"volume", "employees"}:
            return f"{value:,.0f}"

        if isinstance(value, float):
            return f"{value:,.2f}"

        return f"{value:,}"

    return str(value)


def is_long_text_key(key: Any) -> bool:
    normalized = (
        str(key)
        .lower()
        .replace("_", "")
        .replace("-", "")
    )

    return normalized in {
        item.replace("_", "")
        for item in LONG_TEXT_KEYS
    }


def get_item_title(item: dict[str, Any]) -> str | None:
    title_keys = (
        "title",
        "companyName",
        "name",
        "form",
        "symbol",
        "ticker",
    )

    for key in title_keys:
        value = item.get(key)

        if value:
            return format_value(key, value)

    return None


def report_payload(report: Report) -> dict[str, Any]:
    return {
        "id": report.id,
        "ticker": report.ticker,
        "company_name": report.company_name,
        "title": report.title,
        "report_type": report.report_type,
        "created_at": report.created_at.isoformat(),
        "report_data": report.report_data or {},
    }


def safe_filename(ticker: str) -> str:
    return "".join(
        character
        if character.isalnum() or character in ("-", "_")
        else "_"
        for character in ticker
    )


def pdf_text(
    text: Any,
    style: ParagraphStyle,
) -> Paragraph:
    formatted_text = escape(str(text)).replace(
        "\n",
        "<br/>",
    )

    return Paragraph(formatted_text, style)


def build_pdf_value(
    value: Any,
    styles: dict[str, ParagraphStyle],
    skip_keys: set[str] | None = None,
) -> list[Any]:
    if value is None:
        return [
            Paragraph(
                "Not available",
                styles["BodyTextCustom"],
            )
        ]

    if isinstance(value, dict):
        rows: list[list[Any]] = []
        complex_values: list[tuple[str, Any]] = []

        for key, nested_value in value.items():
            if skip_keys and key in skip_keys:
                continue

            if (
                isinstance(nested_value, (dict, list))
                or is_long_text_key(key)
            ):
                complex_values.append((key, nested_value))
            else:
                rows.append(
                    [
                        pdf_text(
                            humanize_key(key),
                            styles["TableLabel"],
                        ),
                        pdf_text(
                            format_value(key, nested_value),
                            styles["BodyTextCustom"],
                        ),
                    ]
                )

        elements: list[Any] = []

        if rows:
            table = Table(
                rows,
                colWidths=[
                    1.8 * inch,
                    5.1 * inch,
                ],
                hAlign="LEFT",
            )

            table.setStyle(
                TableStyle(
                    [
                        (
                            "BACKGROUND",
                            (0, 0),
                            (0, -1),
                            colors.HexColor("#F1F5F9"),
                        ),
                        (
                            "BOX",
                            (0, 0),
                            (-1, -1),
                            0.5,
                            colors.HexColor("#CBD5E1"),
                        ),
                        (
                            "INNERGRID",
                            (0, 0),
                            (-1, -1),
                            0.25,
                            colors.HexColor("#E2E8F0"),
                        ),
                        (
                            "VALIGN",
                            (0, 0),
                            (-1, -1),
                            "TOP",
                        ),
                        (
                            "LEFTPADDING",
                            (0, 0),
                            (-1, -1),
                            7,
                        ),
                        (
                            "RIGHTPADDING",
                            (0, 0),
                            (-1, -1),
                            7,
                        ),
                        (
                            "TOPPADDING",
                            (0, 0),
                            (-1, -1),
                            6,
                        ),
                        (
                            "BOTTOMPADDING",
                            (0, 0),
                            (-1, -1),
                            6,
                        ),
                    ]
                )
            )

            elements.append(table)
            elements.append(Spacer(1, 8))

        for key, nested_value in complex_values:
            elements.append(
                Paragraph(
                    humanize_key(key),
                    styles["Subsection"],
                )
            )

            if (
                is_long_text_key(key)
                and not isinstance(
                    nested_value,
                    (dict, list),
                )
            ):
                elements.append(
                    pdf_text(
                        format_value(key, nested_value),
                        styles["BodyTextCustom"],
                    )
                )
            else:
                elements.extend(
                    build_pdf_value(
                        nested_value,
                        styles,
                    )
                )

            elements.append(Spacer(1, 6))

        return elements

    if isinstance(value, list):
        elements = []

        if not value:
            return [
                Paragraph(
                    "No data available.",
                    styles["BodyTextCustom"],
                )
            ]

        for index, item in enumerate(value, start=1):
            if isinstance(item, dict):
                item_title = get_item_title(item)
                heading = item_title or f"Item {index}"

                elements.append(
                    Paragraph(
                        escape(heading),
                        styles["Subsection"],
                    )
                )

                title_key = next(
                    (
                        key
                        for key in (
                            "title",
                            "companyName",
                            "name",
                            "form",
                            "symbol",
                            "ticker",
                        )
                        if item.get(key)
                    ),
                    None,
                )

                elements.extend(
                    build_pdf_value(
                        item,
                        styles,
                        skip_keys={title_key}
                        if title_key
                        else None,
                    )
                )
            else:
                elements.append(
                    pdf_text(
                        f"• {format_value('', item)}",
                        styles["BodyTextCustom"],
                    )
                )

            elements.append(Spacer(1, 6))

        return elements

    return [
        pdf_text(
            format_value("", value),
            styles["BodyTextCustom"],
        )
    ]


def build_pdf(report: Report) -> BytesIO:
    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.65 * inch,
        leftMargin=0.65 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.65 * inch,
    )

    base_styles = getSampleStyleSheet()

    styles = {
        "Title": ParagraphStyle(
            "ReportTitle",
            parent=base_styles["Title"],
            fontSize=21,
            leading=25,
            spaceAfter=8,
        ),
        "Metadata": ParagraphStyle(
            "ReportMetadata",
            parent=base_styles["Normal"],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#64748B"),
            spaceAfter=3,
        ),
        "Section": ParagraphStyle(
            "ReportSection",
            parent=base_styles["Heading2"],
            fontSize=15,
            leading=19,
            textColor=colors.HexColor("#111827"),
            spaceBefore=16,
            spaceAfter=8,
        ),
        "Subsection": ParagraphStyle(
            "ReportSubsection",
            parent=base_styles["Heading3"],
            fontSize=11,
            leading=15,
            textColor=colors.HexColor("#374151"),
            spaceBefore=7,
            spaceAfter=5,
        ),
        "BodyTextCustom": ParagraphStyle(
            "ReportBody",
            parent=base_styles["BodyText"],
            fontSize=9.5,
            leading=13,
            spaceAfter=4,
        ),
        "TableLabel": ParagraphStyle(
            "ReportTableLabel",
            parent=base_styles["BodyText"],
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#334155"),
        ),
    }

    data = report.report_data or {}

    story: list[Any] = [
        Paragraph(
            escape(report.title),
            styles["Title"],
        ),
        Paragraph(
            escape(
                f"This report summarizes the available "
                f"equity research for {report.company_name} "
                f"({report.ticker})."
            ),
            styles["BodyTextCustom"],
        ),
        Paragraph(
            escape(f"Report type: {report.report_type}"),
            styles["Metadata"],
        ),
        Paragraph(
            escape(
                f"Created: "
                f"{report.created_at.strftime('%B %d, %Y')}"
            ),
            styles["Metadata"],
        ),
        Spacer(1, 12),
    ]

    for section_title, data_key in REPORT_SECTIONS:
        story.append(
            Paragraph(
                escape(section_title),
                styles["Section"],
            )
        )

        story.extend(
            build_pdf_value(
                data.get(data_key),
                styles,
            )
        )

    document.build(story)
    buffer.seek(0)

    return buffer


def add_docx_value(
    document: Document,
    value: Any,
    skip_keys: set[str] | None = None,
) -> None:
    if value is None:
        document.add_paragraph("Not available.")
        return

    if isinstance(value, dict):
        scalar_values: list[tuple[str, Any]] = []
        complex_values: list[tuple[str, Any]] = []

        for key, nested_value in value.items():
            if skip_keys and key in skip_keys:
                continue

            if (
                isinstance(nested_value, (dict, list))
                or is_long_text_key(key)
            ):
                complex_values.append((key, nested_value))
            else:
                scalar_values.append((key, nested_value))

        if scalar_values:
            table = document.add_table(
                rows=1,
                cols=2,
            )
            table.style = "Table Grid"

            header = table.rows[0].cells
            header[0].text = "Field"
            header[1].text = "Value"

            for key, nested_value in scalar_values:
                cells = table.add_row().cells
                cells[0].text = humanize_key(key)
                cells[1].text = format_value(
                    key,
                    nested_value,
                )

        for key, nested_value in complex_values:
            document.add_heading(
                humanize_key(key),
                level=2,
            )

            if (
                is_long_text_key(key)
                and not isinstance(
                    nested_value,
                    (dict, list),
                )
            ):
                document.add_paragraph(
                    format_value(key, nested_value)
                )
            else:
                add_docx_value(
                    document,
                    nested_value,
                )

        return

    if isinstance(value, list):
        if not value:
            document.add_paragraph("No data available.")
            return

        for index, item in enumerate(value, start=1):
            if isinstance(item, dict):
                item_title = get_item_title(item)

                document.add_heading(
                    item_title or f"Item {index}",
                    level=2,
                )

                title_key = next(
                    (
                        key
                        for key in (
                            "title",
                            "companyName",
                            "name",
                            "form",
                            "symbol",
                            "ticker",
                        )
                        if item.get(key)
                    ),
                    None,
                )

                add_docx_value(
                    document,
                    item,
                    skip_keys={title_key}
                    if title_key
                    else None,
                )
            else:
                document.add_paragraph(
                    format_value("", item),
                    style="List Bullet",
                )

        return

    document.add_paragraph(
        format_value("", value)
    )


def build_docx(report: Report) -> BytesIO:
    document = Document()

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10)

    document.add_heading(report.title, level=0)

    document.add_paragraph(
        f"This report summarizes the available equity research "
        f"for {report.company_name} ({report.ticker})."
    )

    document.add_paragraph(
        f"Report type: {report.report_type}"
    )

    document.add_paragraph(
        f"Created: "
        f"{report.created_at.strftime('%B %d, %Y')}"
    )

    data = report.report_data or {}

    for section_title, data_key in REPORT_SECTIONS:
        document.add_heading(section_title, level=1)

        add_docx_value(
            document,
            data.get(data_key),
        )

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    return buffer


def build_json(report: Report) -> BytesIO:
    buffer = BytesIO(
        json.dumps(
            report_payload(report),
            indent=2,
            ensure_ascii=False,
            default=str,
        ).encode("utf-8")
    )

    buffer.seek(0)

    return buffer


def build_report_export(
    report: Report,
    file_format: str,
) -> tuple[BytesIO, str, str]:
    normalized_format = file_format.lower().strip()
    filename_base = (
        f"{safe_filename(report.ticker)}_research_report"
    )

    if normalized_format == "pdf":
        return (
            build_pdf(report),
            "application/pdf",
            f"{filename_base}.pdf",
        )

    if normalized_format == "docx":
        return (
            build_docx(report),
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document",
            f"{filename_base}.docx",
        )

    if normalized_format == "json":
        return (
            build_json(report),
            "application/json",
            f"{filename_base}.json",
        )

    raise ValueError(
        "Unsupported format. Choose pdf, docx, or json."
    )